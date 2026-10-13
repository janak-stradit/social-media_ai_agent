"""Bulk invitations: reading CSV / Excel contact lists, saving them, and
inviting a selection. SMTP is faked - no email is sent."""

import csv
import io
import uuid

import pytest
from openpyxl import Workbook

import db
from services import invite_import as ii


def _csv(rows) -> bytes:
    out = io.StringIO()
    csv.writer(out).writerows(rows)
    return out.getvalue().encode("utf-8")


def _xlsx(rows) -> bytes:
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _email():
    return f"{uuid.uuid4().hex[:10]}@example.com"


# ── Reading files ───────────────────────────────────────────────────────────


def test_csv_with_any_header_spelling_and_order():
    data = _csv([["Company", "E-mail Address", "Full Name", "Designation", "Mobile", "Market", "Remarks"],
                 ["Acme", " Priya@Example.com ", "Priya Shah", "Marketing Manager", "9876543210", "India", "Expo lead"]])
    result = ii.parse("leads.csv", data)
    [row] = result["rows"]
    assert row == {"email": "priya@example.com", "name": "Priya Shah", "company": "Acme", "job_title": "Marketing Manager",
                   "phone": "9876543210", "city": "", "country": "India", "industry": "", "company_size": "", "website": "",
                   "description": "", "notes": "Expo lead", "extra": {}, "shared_inbox": False, "line": 2, "status": "ready"}
    assert result["summary"] == {"ready": 1, "invalid_email": 0, "duplicate": 0, "member": 0, "total": 1, "shared_inbox": 0}
    assert "email" in result["columns"] and "job title" in result["columns"]


def test_excel_files_are_read_too():
    data = _xlsx([["email", "First name", "Last name", "Phone"], ["ravi@example.com", "Ravi", "Kumar", 9876543210], [None, None, None, None]])
    [row] = ii.parse("contacts.xlsx", data)["rows"]  # the blank row is skipped
    assert row["name"] == "Ravi Kumar" and row["phone"] == "9876543210"  # not "9876543210.0"


def test_every_row_is_checked():
    data = _csv([["email", "name"], ["a@example.com", "A"], ["not-an-email", "B"], ["", "C"], ["A@example.com", "A again"],
                 ["member@example.com", "D"]])
    result = ii.parse("x.csv", data, existing_member_emails={"member@example.com"})
    assert [r["status"] for r in result["rows"]] == ["ready", "invalid_email", "invalid_email", "duplicate", "member"]
    assert [r["line"] for r in result["rows"]] == [2, 3, 4, 5, 6]
    assert result["summary"] == {"ready": 1, "invalid_email": 2, "duplicate": 1, "member": 1, "total": 5, "shared_inbox": 0}


def test_semicolon_csv_and_a_list_without_a_header():
    assert ii.parse("x.csv", b"email;name\nsam@example.com;Sam\n")["rows"][0]["name"] == "Sam"
    rows = ii.parse("x.csv", b"one@example.com\ntwo@example.com\n")["rows"]
    assert [r["email"] for r in rows] == ["one@example.com", "two@example.com"]


@pytest.mark.parametrize("filename, data, message", [
    ("x.csv", b"", "empty"),
    ("x.csv", b"name,company\nSam,Acme\n", "No email column"),
    ("x.pdf", b"%PDF", ".csv or .xlsx"),
    ("x.xls", b"\xd0\xcf", "Old .xls"),
    ("x.xlsx", b"not a workbook", "could not be read"),
    ("x.csv", b"x" * (ii.MAX_FILE_BYTES + 1), "2 MB"),
], ids=["empty", "no-email-column", "pdf", "old-xls", "broken-xlsx", "too-big"])
def test_unreadable_files_are_explained(filename, data, message):
    with pytest.raises(ii.InviteFileError, match=message):
        ii.parse(filename, data)


def test_too_many_rows():
    rows = [["email"]] + [[f"p{i}@example.com"] for i in range(ii.MAX_ROWS + 1)]
    with pytest.raises(ii.InviteFileError, match="at most 2,000"):
        ii.parse("x.csv", _csv(rows))


LEAD_SHEET = ["City", "Tier", "Size", "Domain/Industry", "Company Name", "Contact Person (if published)", "Designation",
              "Business Email", "Website", "Business Description (1–2 lines)", "Source URL(s)", "Confidence", "Date Collected",
              "Size Source", "Notes / Flags", "Dedupe Key (domain)"]


def test_a_lead_sheet_is_stored_field_by_field_and_other_columns_are_kept():
    import datetime

    book = Workbook()
    book.active.title = "Summary"  # a summary sheet first: the sheet with the emails is found
    book.active.append(["Batch", "Pune 1"])
    leads = book.create_sheet("Leads")
    leads.append(LEAD_SHEET)
    leads.append(["Pune", "Tier 1", "Small", "IT services", "Synkrama Technologies", "", "", "contact@synkrama.com",
                  "https://www.synkrama.com", "Custom software.", "https://a.example ; https://b.example", "Medium",
                  datetime.datetime(2026, 10, 12), "10-49", "Generic inbox", "synkrama.com"])  # noqa: DTZ001 - Excel dates are naive
    leads.append(["Pune", "Tier 2", "Mid", "Manufacturing", "Beta Works", "Asha Rao", "Head of Marketing", "asha.rao@beta.example",
                  "beta.example", "", "", "High", "2026-10-12", "", "", "beta.example"])
    buf = io.BytesIO()
    book.save(buf)
    result = ii.parse("Pune_Leads.xlsx", buf.getvalue())
    first, second = result["rows"]
    assert (first["company"], first["city"], first["country"], first["industry"], first["company_size"], first["website"],
            first["description"], first["notes"]) == ("Synkrama Technologies", "Pune", "", "IT services", "Small",
                                                     "https://www.synkrama.com", "Custom software.", "Generic inbox")
    assert first["extra"] == {"Tier": "Tier 1", "Source URL(s)": "https://a.example ; https://b.example", "Confidence": "Medium",
                              "Date Collected": "2026-10-12", "Size Source": "10-49", "Dedupe Key (domain)": "synkrama.com"}
    assert first["shared_inbox"] and not second["shared_inbox"]
    assert (second["name"], second["job_title"]) == ("Asha Rao", "Head of Marketing")
    assert result["summary"]["ready"] == 2 and result["summary"]["shared_inbox"] == 1
    assert result["extra_columns"] == ["Tier", "Source URL(s)", "Confidence", "Date Collected", "Size Source", "Dedupe Key (domain)"]

    # Saved with the extra details; a later upload adds to them rather than replacing them
    admin = db.create_user("Admin", _email(), "h", is_admin=True)
    email = _email()
    rows = ii.clean_rows([{**first, "email": email}])
    db.save_invite_contacts(rows, "Pune_Leads.xlsx", admin["id"])
    db.save_invite_contacts(ii.clean_rows([{"email": email, "extra": {"Confidence": "High", "Owner": "Janak"}}]), "b.csv", admin["id"])
    saved = next(c for c in db.list_invite_contacts()["contacts"] if c["email"] == email)
    assert saved["industry"] == "IT services" and saved["website"] == "https://www.synkrama.com" and saved["city"] == "Pune"
    assert saved["extra"]["Tier"] == "Tier 1" and saved["extra"]["Confidence"] == "High" and saved["extra"]["Owner"] == "Janak"


def test_template_has_the_expected_columns():
    header = next(csv.reader(io.StringIO(ii.template_csv())))
    assert header == ["email", "name", "company", "job_title", "phone", "city", "country", "industry", "company_size", "website",
                      "description", "notes"]
    assert ii.parse("t.csv", ii.template_csv().encode())["summary"]["ready"] == 1


# ── Saving ──────────────────────────────────────────────────────────────────


def test_saving_adds_new_people_and_updates_known_ones_without_wiping():
    admin = db.create_user("Admin", _email(), "h", is_admin=True)
    e1, e2 = _email(), _email()
    first = db.save_invite_contacts([{"email": e1, "name": "Priya", "company": "Acme", "phone": "123"}], "a.csv", admin["id"], rows_total=3)
    assert first["added"] == 1 and first["updated"] == 0 and first["skipped"] == 2
    second = db.save_invite_contacts([{"email": e1, "name": "", "company": "Acme Ltd"}, {"email": e2, "name": "Ravi"}], "b.xlsx", admin["id"])
    assert second["added"] == 1 and second["updated"] == 1
    contacts = {c["email"]: c for c in db.list_invite_contacts()["contacts"]}
    assert contacts[e1]["name"] == "Priya" and contacts[e1]["company"] == "Acme Ltd" and contacts[e1]["phone"] == "123"
    assert contacts[e1]["batch"] == "b.xlsx" and contacts[e1]["status"] == "not_invited"
    assert contacts[e1]["shared_inbox"] is False
    db.save_invite_contacts([{"email": "sales@" + e2.split("@")[1]}], None, admin["id"], source="manual")
    assert {c["email"]: c for c in db.list_invite_contacts()["contacts"]}["sales@" + e2.split("@")[1]]["shared_inbox"] is True


# ── Through the admin page ──────────────────────────────────────────────────


@pytest.fixture
def admin_client():
    from app import create_app

    app = create_app("development")
    app.config.update(TESTING=True)
    admin = db.create_user("Admin", _email(), "h", is_admin=True)
    client = app.test_client()
    with client.session_transaction() as s:
        s["user_id"] = admin["id"]
    return client


@pytest.fixture
def mail(monkeypatch):
    from api import routes

    sent, fail = [], set()

    def send(invitation, token, *_args):
        if invitation["email"] in fail:
            raise RuntimeError("Mailbox unavailable")
        sent.append((invitation["email"], token))

    monkeypatch.setattr(routes, "_send_invitation", send)
    return sent, fail


def test_upload_preview_import_and_send(admin_client, mail):
    sent, fail = mail
    member = db.create_user("Existing", _email(), "h")["email"]
    good, bad = _email(), _email()
    data = _xlsx([["Email", "Name", "Company"], [good, "Priya", "Acme"], [bad, "Ravi", "Beta"], [member, "Old", "X"], ["nope", "Z", ""]])

    preview = admin_client.post("/api/admin/invite-contacts/preview", data={"file": (io.BytesIO(data), "leads.xlsx")},
                                content_type="multipart/form-data").get_json()
    assert preview["success"] and preview["summary"] == {"ready": 2, "invalid_email": 1, "duplicate": 0, "member": 1, "total": 4,
                                                         "shared_inbox": 0}
    assert db.list_invite_contacts()["counts"]["all"] >= 0  # nothing saved by a preview
    assert not [c for c in db.list_invite_contacts()["contacts"] if c["email"] == good]

    imported = admin_client.post("/api/admin/invite-contacts/import", json={
        "filename": preview["filename"], "total": 4, "rows": preview["rows"]}).get_json()  # the page sends every row; only the good ones are kept
    assert imported["added"] == 2 and set(imported["emails"]) == {good, bad}
    ids = {c["email"]: c["id"] for c in imported["contacts"]}
    assert member not in ids

    fail.add(bad)
    result = admin_client.post("/api/admin/invite-contacts/send", json={"ids": [ids[good], ids[bad]], "inviter_name": "Janak", "message": "Welcome!"}).get_json()
    assert result["invited"] == 1 and result["failed"] == 1
    assert [e for e, _ in sent] == [good]
    statuses = {c["email"]: c for c in admin_client.get("/api/admin/invite-contacts").get_json()["contacts"]}
    assert statuses[good]["status"] == "invited" and statuses[good]["invite_count"] == 1 and statuses[good]["last_invited_at"]
    assert statuses[bad]["status"] == "failed" and "Mailbox unavailable" in statuses[bad]["last_error"]
    invitation = next(i for i in db.list_invitations() if i["email"] == good)
    assert invitation["invited_by"] == "Janak" and invitation["message"] == "Welcome!"

    # The failed one is retried and now goes through
    fail.clear()
    retry = admin_client.post("/api/admin/invite-contacts/send", json={"ids": [ids[bad]]}).get_json()
    assert retry["invited"] == 1
    assert {c["email"]: c["status"] for c in admin_client.get("/api/admin/invite-contacts").get_json()["contacts"]}[bad] == "invited"


def test_someone_who_signed_up_shows_as_joined_and_is_skipped(admin_client, mail):
    sent, _ = mail
    email = _email()
    contact_id = admin_client.post("/api/admin/invite-contacts", json={"email": email, "name": "Sam"}).get_json()["contacts"][0]["id"]
    db.create_user("Sam", email, "h")
    contact = next(c for c in admin_client.get("/api/admin/invite-contacts").get_json()["contacts"] if c["id"] == contact_id)
    assert contact["status"] == "joined" and contact["source"] == "manual"
    result = admin_client.post("/api/admin/invite-contacts/send", json={"ids": [contact_id]}).get_json()
    assert result["joined"] == 1 and not sent


def test_limits_and_removal(admin_client, mail):
    assert admin_client.post("/api/admin/invite-contacts/send", json={"ids": []}).status_code == 400
    assert admin_client.post("/api/admin/invite-contacts/send", json={"ids": list(range(1, 27))}).status_code == 400
    assert admin_client.post("/api/admin/invite-contacts", json={"email": "bad"}).status_code == 400
    email = _email()
    contact_id = next(c["id"] for c in admin_client.post("/api/admin/invite-contacts", json={"email": email}).get_json()["contacts"] if c["email"] == email)
    r = admin_client.delete("/api/admin/invite-contacts", json={"ids": [contact_id]}).get_json()
    assert r["removed"] == 1 and not [c for c in r["contacts"] if c["email"] == email]


def test_template_and_export_downloads(admin_client):
    template = admin_client.get("/api/admin/invite-contacts/template.csv")
    assert template.status_code == 200 and template.headers["Content-Type"].startswith("text/csv")
    assert template.get_data(as_text=True).startswith("email,name,company")
    email = _email()
    admin_client.post("/api/admin/invite-contacts", json={"email": email, "company": "Acme"})
    export = admin_client.get("/api/admin/invite-contacts/export.csv").get_data(as_text=True)
    assert export.startswith("email,name,company") and email in export and "not_invited" in export


def test_only_admins(admin_client):
    from app import create_app

    app = create_app("development")
    app.config.update(TESTING=True)
    user = db.create_user("User", _email(), "h")
    client = app.test_client()
    with client.session_transaction() as s:
        s["user_id"] = user["id"]
    for method, url in (("get", "/api/admin/invite-contacts"), ("post", "/api/admin/invite-contacts/send"),
                        ("get", "/api/admin/invite-contacts/export.csv"), ("post", "/api/admin/invite-contacts/preview")):
        assert getattr(client, method)(url).status_code in (401, 403)
    assert app.test_client().get("/api/admin/invite-contacts").status_code == 401
