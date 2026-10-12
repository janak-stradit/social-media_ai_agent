"""Bulk invitations (Admin -> Invitations -> Contacts): reading a CSV or
Excel file of people into the invite contact list.

The file needs an email column; the other columns are optional and matched
by their header, in any order and with common spellings ("E-mail",
"Full name", "Organisation", "Designation", ...). Every row is checked and
given a status before anything is saved:

  ready           a new contact, or an update to one already on the list
  invalid_email   the email isn't a valid address (or is missing)
  duplicate       the same email appears earlier in the file
  member          that email already has an AVIR AI account

The admin sees this preview and imports the ready rows; members are kept
out of the list (they don't need an invitation).
"""

import csv
import datetime
import io
import re

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# The stored columns: (field, header in the template, max length). Any other
# column in the file is kept too, under "extra" ({header: value}).
FIELDS = [
    ("email", "email", 255),
    ("name", "name", 255),
    ("company", "company", 255),
    ("job_title", "job_title", 255),
    ("phone", "phone", 40),
    ("city", "city", 80),
    ("country", "country", 80),
    ("industry", "industry", 120),
    ("company_size", "company_size", 40),
    ("website", "website", 255),
    ("description", "description", 1000),
    ("notes", "notes", 1000),
]
MAX_EXTRA_COLUMNS, MAX_EXTRA_VALUE = 20, 500
LIMITS = {field: limit for field, _, limit in FIELDS}

# Header spellings understood for each field (compared lower-case, without spaces, _ or -)
ALIASES = {
    "email": ("email", "emailaddress", "mail", "workemail", "businessemail", "emailid", "officialemail"),
    "name": ("name", "fullname", "contactname", "contact", "person", "contactperson", "keycontact", "decisionmaker"),
    "first_name": ("firstname", "givenname", "first"),
    "last_name": ("lastname", "surname", "familyname", "last"),
    "company": ("company", "companyname", "organisation", "organization", "organisationname", "organizationname", "business",
                "businessname", "brand", "account", "accountname"),
    "job_title": ("jobtitle", "title", "role", "designation", "position"),
    "phone": ("phone", "phonenumber", "mobile", "mobilenumber", "telephone", "tel", "contactnumber", "whatsapp"),
    "city": ("city", "town", "location"),
    "country": ("country", "market", "region"),
    "industry": ("industry", "domainindustry", "sector", "vertical", "domain", "category", "segment"),
    "company_size": ("size", "companysize", "employees", "employeecount", "headcount", "teamsize"),
    "website": ("website", "site", "url", "web", "websiteurl", "homepage"),
    "description": ("description", "businessdescription", "about", "companydescription", "summary", "profile"),
    "notes": ("notes", "note", "comments", "comment", "remarks", "notesflags", "flags"),
}
_LOOKUP = {alias: field for field, aliases in ALIASES.items() for alias in aliases}
# Addresses that reach a team's shared inbox rather than one person (flagged in the preview)
SHARED_INBOXES = {"info", "contact", "contactus", "sales", "hello", "hi", "enquiry", "enquiries", "inquiry", "inquiries", "support",
                  "admin", "office", "mail", "team", "marketing", "careers", "hr", "business", "connect", "help"}


class InviteFileError(ValueError):
    """The file can't be read (shown to the admin)."""


def _key(header) -> str:
    """'Contact Person (if published)' -> 'contactperson', 'E-mail' -> 'email'."""
    text = re.sub(r"\(.*?\)", "", str(header or "")).strip().lower()
    return re.sub(r"[^a-z0-9]", "", text)


def _clean(value, limit: int) -> str:
    if isinstance(value, datetime.datetime) and value.time() == datetime.time():
        value = value.date()  # an Excel date cell: "2026-10-12", not "2026-10-12 00:00:00"
    text = "" if value is None else str(value)
    if text.endswith(".0") and text[:-2].isdigit():  # Excel turns 9876543210 into 9876543210.0
        text = text[:-2]
    return " ".join(text.split())[:limit]


def template_csv() -> str:
    """The file admins can download, fill in and upload."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([header for _, header, _ in FIELDS])
    writer.writerow(["priya.shah@example.com", "Priya Shah", "Acme Logistics", "Marketing Manager", "+91 98765 43210", "Pune", "India",
                     "Logistics", "11-50", "https://acme-logistics.example", "Freight and warehousing across Maharashtra.",
                     "Met at the Mumbai expo"])
    return out.getvalue()


def _read_table(filename: str, data: bytes) -> list[list]:
    name = (filename or "").lower()
    if not data:
        raise InviteFileError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise InviteFileError("The file is larger than 2 MB. Split it into smaller files.")
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = [[list(r) for r in sheet.iter_rows(values_only=True)] for sheet in book.worksheets]
            book.close()
            for rows in sheets:
                header = next((r for r in rows if any(c not in (None, "") for c in r)), [])
                if any(_LOOKUP.get(_key(h)) == "email" for h in header):
                    return rows
            return sheets[0] if sheets else []
        except Exception as err:  # a broken or protected workbook
            raise InviteFileError(f"The Excel file could not be read ({type(err).__name__}). Save it again as .xlsx or .csv.") from err
    if name.endswith(".xls"):
        raise InviteFileError("Old .xls files aren't supported. Save the file as .xlsx or .csv and upload it again.")
    if not name.endswith((".csv", ".txt")):
        raise InviteFileError("Upload a .csv or .xlsx file.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def parse(filename: str, data: bytes, existing_member_emails=None) -> dict:
    """{"rows": [...], "summary": {...}, "columns": {...}} - nothing is saved.
    existing_member_emails: lower-case emails that already have accounts."""
    table = [r for r in _read_table(filename, data) if any(str(c or "").strip() for c in r)]
    if not table:
        raise InviteFileError("The file has no rows.")
    headers = [_key(h) for h in table[0]]
    columns = {}
    for index, header in enumerate(headers):
        field = _LOOKUP.get(header)
        if field and field not in columns:
            columns[field] = index
    used = set(columns.values())
    extra_columns = [(i, " ".join(str(h).split())[:60]) for i, h in enumerate(table[0])
                     if i not in used and str(h or "").strip()][:MAX_EXTRA_COLUMNS]
    if "email" not in columns:
        # No header row: a column of email addresses still works
        first = [str(c or "").strip() for c in table[0]]
        email_col = next((i for i, c in enumerate(first) if EMAIL_RE.match(c)), None)
        if email_col is None:
            raise InviteFileError("No email column found. The first row should name the columns, e.g. email, name, company.")
        columns, body, start_line, extra_columns = {"email": email_col}, table, 1, []
    else:
        body, start_line = table[1:], 2
    if len(body) > MAX_ROWS:
        raise InviteFileError(f"The file has {len(body):,} rows. Upload at most {MAX_ROWS:,} at a time.")

    members = {e.lower() for e in (existing_member_emails or [])}
    seen, rows = set(), []
    for offset, raw in enumerate(body):
        def cell(field, raw=raw):
            i = columns.get(field)
            return raw[i] if i is not None and i < len(raw) else None

        row = {field: _clean(cell(field), LIMITS[field]) for field in LIMITS}
        if not row["name"]:
            row["name"] = _clean(" ".join(str(cell(f) or "") for f in ("first_name", "last_name")), LIMITS["name"])
        row["email"] = row["email"].lower().strip("<>").replace("mailto:", "")
        row["extra"] = {name: _clean(raw[i], MAX_EXTRA_VALUE) for i, name in extra_columns
                        if i < len(raw) and _clean(raw[i], MAX_EXTRA_VALUE)}
        row["shared_inbox"] = row["email"].split("@")[0] in SHARED_INBOXES
        row["line"] = start_line + offset
        if not EMAIL_RE.match(row["email"]):
            row["status"] = "invalid_email"
        elif row["email"] in seen:
            row["status"] = "duplicate"
        elif row["email"] in members:
            row["status"] = "member"
        else:
            row["status"] = "ready"
        if EMAIL_RE.match(row["email"]):
            seen.add(row["email"])
        rows.append(row)

    summary = {status: sum(1 for r in rows if r["status"] == status) for status in ("ready", "invalid_email", "duplicate", "member")}
    summary["total"] = len(rows)
    summary["shared_inbox"] = sum(1 for r in rows if r["status"] == "ready" and r["shared_inbox"])
    names = {"first_name": "first name", "last_name": "last name"}
    return {"rows": rows, "summary": summary,
            "columns": sorted(names.get(f, f.replace("_", " ")) for f in columns),
            "extra_columns": [name for _, name in extra_columns]}


def clean_rows(raw_rows) -> list[dict]:
    """Rows sent back by the page for import: only the stored fields, cleaned again."""
    out, seen = [], set()
    for raw in raw_rows or []:
        if not isinstance(raw, dict):
            continue
        row = {field: _clean(raw.get(field), LIMITS[field]) for field in LIMITS}
        row["email"] = row["email"].lower()
        extra = raw.get("extra") if isinstance(raw.get("extra"), dict) else {}
        row["extra"] = {_clean(k, 60): _clean(v, MAX_EXTRA_VALUE) for k, v in list(extra.items())[:MAX_EXTRA_COLUMNS]
                        if _clean(k, 60) and _clean(v, MAX_EXTRA_VALUE)}
        if EMAIL_RE.match(row["email"]) and row["email"] not in seen:
            seen.add(row["email"])
            out.append(row)
    return out[:MAX_ROWS]
